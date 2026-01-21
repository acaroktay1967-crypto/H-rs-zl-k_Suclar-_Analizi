"""
OpenXML Document Generator
Generates Word documents from analysis results using python-docx
"""

from typing import Dict
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


class OpenXMLGenerator:
    """Generator for OpenXML (Word document) formatted analysis reports"""
    
    def generate_report(self, analysis_data: Dict, output_path: str) -> bool:
        """
        Generate Word document report from analysis data
        
        Args:
            analysis_data: Analysis results dictionary
            output_path: Path to save the document
            
        Returns:
            True if successful, False otherwise
        """
        try:
            doc = Document()
            
            # Set document properties
            doc.core_properties.title = "Hırsızlık Suçları Analiz Raporu"
            doc.core_properties.author = "AI Legal Analysis System"
            doc.core_properties.created = datetime.now()
            
            # Title
            title = doc.add_heading('Hırsızlık Suçları Analiz Raporu', 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Metadata
            metadata = doc.add_paragraph()
            metadata.add_run(f"Rapor Tarihi: {datetime.now().strftime('%d.%m.%Y %H:%M')}").italic = True
            metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            doc.add_paragraph()  # Spacer
            
            # Event Information Section
            doc.add_heading('Olay Bilgileri', 1)
            
            event_type = doc.add_paragraph()
            event_type.add_run('Olay Türü: ').bold = True
            event_type.add_run(analysis_data.get('event_type', 'Belirtilmemiş'))
            
            event_desc = doc.add_paragraph()
            event_desc.add_run('Olay Açıklaması: ').bold = True
            event_desc.add_run(analysis_data.get('event_description', ''))
            
            # Keywords
            keywords = analysis_data.get('keywords', [])
            if keywords:
                kw_para = doc.add_paragraph()
                kw_para.add_run('Anahtar Kelimeler: ').bold = True
                kw_para.add_run(', '.join(keywords))
            
            doc.add_paragraph()  # Spacer
            
            # OpenAI Analysis Section
            openai_data = analysis_data.get('openai_analysis', {})
            if openai_data.get('success'):
                doc.add_heading('OpenAI Analizi', 1)
                doc.add_paragraph(openai_data.get('analysis', ''))
                
                # Model info
                model_info = doc.add_paragraph()
                model_info.add_run(f"Model: {openai_data.get('model', '')}").italic = True
                
                doc.add_paragraph()  # Spacer
            
            # Recommendations Section
            recommendations = analysis_data.get('recommendations', {})
            if recommendations.get('success'):
                doc.add_heading('Hukuki Öneriler', 1)
                doc.add_paragraph(recommendations.get('recommendations', ''))
                doc.add_paragraph()  # Spacer
            
            # Yargıtay Decisions Section
            yargitay = analysis_data.get('yargitay_decisions', {})
            if yargitay.get('success'):
                doc.add_page_break()
                doc.add_heading('Yargıtay Kararları', 1)
                doc.add_paragraph(yargitay.get('decisions', ''))
                doc.add_paragraph()  # Spacer
            
            # Correlation Analysis Section
            correlation = analysis_data.get('correlation_analysis', {})
            if correlation.get('success'):
                doc.add_heading('İlişkilendirme Analizi', 1)
                doc.add_paragraph(correlation.get('correlation_analysis', ''))
                doc.add_paragraph()  # Spacer
            
            # Academic Sources Section
            academic = analysis_data.get('academic_sources', {})
            if academic.get('success'):
                doc.add_page_break()
                doc.add_heading('Akademik Kaynaklar', 1)
                doc.add_paragraph(academic.get('academic_sources', ''))
            
            # Footer
            doc.add_paragraph()
            footer = doc.add_paragraph()
            footer.add_run(
                f"\nBu rapor yapay zeka destekli analiz sistemi tarafından "
                f"{datetime.now().strftime('%d.%m.%Y')} tarihinde oluşturulmuştur."
            ).italic = True
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Save document
            doc.save(output_path)
            return True
            
        except Exception as e:
            print(f"OpenXML generation error: {e}")
            return False
